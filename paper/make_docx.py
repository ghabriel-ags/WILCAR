"""
Word version of the manuscript in the format of Neural Networks (Elsevier guide for authors):
editable equations (OMML) and tables (Word tables), APA 7th edition references, numbered sections, Times New Roman
12 pt, double spacing, continuous line numbers, page numbers, title page with affiliations and corresponding author,
abstract and keywords; highlights in a separate file.

    python paper/make_docx.py            ->  paper/main.docx, paper/highlights.docx

Requires a previous LaTeX build (main.aux / main.bbl provide the numbers of sections, equations, theorems, tables,
figures and algorithms, so the Word text cites exactly what the PDF shows).
"""
import re, subprocess
from pathlib import Path

HERE = Path(__file__).resolve().parent
NBSP = "\u00a0"

HIGHLIGHTS = [
    "Constructive single-hidden-layer classifiers with certified partial monotonicity",
    "Monotone single-hidden-layer networks are universal; sign-constrained ones are not",
    "Counterexample-guided training returns a network that is always certified",
    "Branch-and-bound certificate, exact in pre-activation space, needs no solver",
    "Nested evaluation against Min-Max, constrained and Lipschitz monotonic networks",
]

TYPES = {"section": "Section", "subsection": "Section", "equation": "Eq.", "algorithm": "Algorithm",
         "theorem": "Theorem", "proposition": "Proposition", "lemma": "Lemma", "definition": "Definition",
         "table": "Table", "figure": "Fig.", "appendix": "", "corollary": "Corollary", "remark": "Remark"}
PLURAL = {"Section": "Sections", "Eq.": "Eqs.", "Algorithm": "Algorithms", "Theorem": "Theorems",
          "Proposition": "Propositions", "Lemma": "Lemmas", "Table": "Tables", "Fig.": "Figs."}


# ------------------------------------------------------------------------------------------ helpers
def braced(s, i):
    """Return (content, index after closing brace) for the brace group starting at s[i] == '{'."""
    assert s[i] == "{", s[i:i + 20]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{" and s[j - 1] != "\\":
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces")


def labels():
    aux = (HERE / "main.aux").read_text()
    out = {}
    for m in re.finditer(r"\\newlabel\{([^}@]+)@cref\}\{\{\[([a-z]+)\]\[[^\]]*\]\[[^\]]*\]([^}]*)\}", aux):
        out[m.group(1)] = (" ".join(m.group(3).replace("~", " ").split()), m.group(2))
    return out


def resolve_refs(tex, lab):
    def one(k):
        return lab.get(k.strip(), ("??", "section"))

    def cref(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        name = TYPES.get(one(keys[0])[1], "")
        if len(keys) > 1:
            name = PLURAL.get(name, name)
        return (name + NBSP if name else "") + " and ".join(one(k)[0] for k in keys)
    tex = re.sub(r"\\[cC]ref\{([^}]+)\}", cref, tex)
    tex = re.sub(r"\\eqref\{([^}]+)\}", lambda m: "(" + one(m.group(1))[0] + ")", tex)
    tex = re.sub(r"\\ref\{([^}]+)\}", lambda m: one(m.group(1))[0], tex)
    return tex


def strip_resizebox(tex):
    out, i = [], 0
    while True:
        j = tex.find("\\resizebox", i)
        if j < 0:
            out.append(tex[i:]); break
        out.append(tex[i:j])
        k = j + len("\\resizebox")
        _, k = braced(tex, k); _, k = braced(tex, k)
        body, k = braced(tex, k)
        out.append(body); i = k
    return "".join(out)


def flatten_inputs(tex):
    def repl(m):
        f = HERE / m.group(1)
        return f.read_text() if f.exists() else ""
    tex = re.sub(r"\\IfFileExists\{(tables/[^}]+)\}\{\\input\{[^}]+\}\}\{\}", repl, tex)
    return tex


def resolve_ifexists(tex):
    """\\IfFileExists{f}{A}{B} -> A if f exists (PDF figures become PNG) else B."""
    out, i = [], 0
    while True:
        j = tex.find("\\IfFileExists", i)
        if j < 0:
            out.append(tex[i:]); break
        out.append(tex[i:j])
        k = j + len("\\IfFileExists")
        f, k = braced(tex, k); a, k = braced(tex, k); b, k = braced(tex, k)
        exists = (HERE / f).exists()
        if exists:
            a = re.sub(r"\{([^{}]+)\.pdf\}", lambda m: "{" + m.group(1) + ".png}"
                       if (HERE / (m.group(1) + ".png")).exists() else m.group(0), a)
        out.append(a if exists else b); i = k
    return "".join(out)


# ------------------------------------------------------------------------------------------ algorithms
KW = {"If": ("if", "then"), "ElsIf": ("else if", "then"), "For": ("for", "do"), "While": ("while", "do")}


def algorithmic_to_text(body):
    """Render an algorithmic block as numbered lines (editable text)."""
    lines, indent, cur = [], 0, None
    i = 0
    toks = re.compile(r"\\(Require|Ensure|State|If|ElsIf|Else|EndIf|For|EndFor|While|EndWhile|Return|Comment)\b")

    def push(text, ind, numbered=True):
        lines.append([text.strip(), ind, numbered])

    while i < len(body):
        m = toks.search(body, i)
        if not m:
            if lines:
                lines[-1][0] += " " + body[i:].strip()
            break
        if lines and body[i:m.start()].strip():
            lines[-1][0] += " " + body[i:m.start()].strip()
        cmd = m.group(1); k = m.end()
        while k < len(body) and body[k] == " ":
            k += 1
        if cmd in ("If", "ElsIf", "For", "While"):
            arg, k = braced(body, k)
            a, b = KW[cmd]
            if cmd == "ElsIf":
                indent -= 1
            push(f"\\textbf{{{a}}} {arg} \\textbf{{{b}}}", indent)
            indent += 1
        elif cmd == "Else":
            push("\\textbf{else}", indent - 1)
        elif cmd in ("EndIf", "EndFor", "EndWhile"):
            indent -= 1
        elif cmd in ("State",):
            push("", indent)
        elif cmd == "Return":
            push("\\textbf{return}", indent)
        elif cmd in ("Require", "Ensure"):
            push("\\textit{Input:}" if cmd == "Require" else "\\textit{Output:}", 0, numbered=False)
        elif cmd == "Comment":
            arg, k = braced(body, k)
            lines[-1][0] += f" \u25b7 \\textit{{{arg}}}"
        i = k
    out, n = [], 0
    for text, ind, numbered in lines:
        if not text:
            continue
        if numbered:
            n += 1
        prefix = (f"{n}:" if numbered else "") + NBSP * 2 + NBSP * 4 * max(ind, 0)
        out.append("\\noindent{}" + prefix + text + "\n")      # '{}': pandoc swallows a number after \\noindent
    return "\n".join(out)


# ------------------------------------------------------------------------------------------ floats
def convert_floats(tex, lab):
    pat = re.compile(r"\\begin\{(algorithm|table\*?|figure\*?)\}(\[[^\]]*\])?(.*?)\\end\{\1\}", re.S)

    def repl(m):
        env, body = m.group(1).rstrip("*"), m.group(3)
        cap_txt, lab_key = "", None
        j = body.find("\\caption{")
        if j >= 0:
            cap_txt, k = braced(body, j + len("\\caption"))
            body = body[:j] + body[k:]
        lm = re.search(r"\\label\{([^}]+)\}", body)
        if lm:
            lab_key = lm.group(1); body = body.replace(lm.group(0), "")
        num = lab.get(lab_key, ("", env))[0] if lab_key else ""
        head = {"algorithm": "Algorithm", "table": "Table", "figure": "Fig."}[env]
        caption = f"\n\n\\noindent\\textbf{{{head} {num}.}} {cap_txt}\n\n"
        body = re.sub(r"\\(centering|small|footnotesize|smallskip)\b", "", body)
        body = body.replace("\\par", "\n\n")
        if env == "algorithm":
            alg = re.search(r"\\begin\{algorithmic\}(\[[^\]]*\])?(.*?)\\end\{algorithmic\}", body, re.S)
            return caption + algorithmic_to_text(alg.group(2)) + "\n\n"
        if env == "table":
            return caption + body + "\n\n"
        return "\n\n" + body + caption      # figure: image first, caption below
    return pat.sub(repl, tex)


def number_headings(tex):
    """Number sections as in the PDF (1., 1.1., Appendix A., A.1.), leaving starred sections unnumbered."""
    out, sec, sub, app = [], 0, 0, False
    for line in tex.split("\n"):
        if line.strip() == "\\appendix":
            app, sec = True, 0
            continue
        m = re.match(r"\\(sub)?section(\*?)\{(.*)\}(\\label\{[^}]+\})?\s*$", line)
        if m and not m.group(2):
            if m.group(1):
                sub += 1
                num = (chr(64 + sec) if app else str(sec)) + f".{sub}."
                line = f"\\subsection*{{{num} {m.group(3)}}}"
            else:
                sec += 1; sub = 0
                num = f"Appendix {chr(64 + sec)}." if app else f"{sec}."
                line = f"\\section*{{{num} {m.group(3)}}}"
        out.append(line)
    return "\n".join(out)


def title_block(tex):
    fm = re.search(r"\\begin\{frontmatter\}(.*?)\\end\{frontmatter\}", tex, re.S).group(1)
    title, _ = braced(fm, fm.index("\\title{") + len("\\title"))
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", fm, re.S).group(1).strip()
    kw = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", fm, re.S).group(1)
    kw = "; ".join(k.strip() for k in kw.split("\\sep"))
    authors = re.findall(r"\\author\[[^\]]*\]\{([^}]*?)(\\corref\{[^}]*\})?\}", fm)
    emails = re.findall(r"\\ead\{([^}]*)\}", fm)
    names = [a + ("\\textsuperscript{a,*}" if c else "\\textsuperscript{a}") for a, c in authors]
    org = re.search(r"organization=\{([^}]*)\}", fm).group(1)
    block = (
        "\\noindent " + ", ".join(names) + "\n\n"
        + "\\noindent \\textsuperscript{a}" + org + ", Salvador, BA, Brazil\n\n"
        + "\\noindent \\textsuperscript{*}Corresponding author. E-mail addresses: "
        + ", ".join(f"{e} ({a[0]})" for e, a in zip(emails, authors)) + "\n\n"
        + "\\section*{Abstract}\n\n" + abstract + "\n\n"
        + "\\noindent\\textit{Keywords:} " + kw + "\n\n\\newpage\n"
    )
    tex = tex.replace(re.search(r"\\begin\{frontmatter\}.*?\\end\{frontmatter\}", tex, re.S).group(0), block)
    return tex, title


# ------------------------------------------------------------------------------------------ styling
def reference_doc(path):
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_LINE_SPACING
    raw = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    path.write_bytes(raw)
    d = docx.Document(str(path))
    for sec in d.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for st in d.styles:
        if st.type != 1:     # paragraph styles only
            continue
        try:
            st.font.name = "Times New Roman"
            rpr = st.element.get_or_add_rPr()
            rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
            if rfonts is not None:
                for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + a, "Times New Roman")
            st.font.color.rgb = None
        except Exception:
            pass
    S = {s.name: s for s in d.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Abstract", "Bibliography", "Block Text"):
        if name in S:
            st = S[name]
            st.font.size = Pt(12)
            pf = st.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            pf.space_before, pf.space_after = Pt(0), Pt(0)
    for name in ("Compact",):
        if name in S:
            st = S[name]; st.font.size = Pt(10)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    for name, size in (("Title", 16), ("Heading 1", 12), ("Heading 2", 12), ("Heading 3", 12)):
        if name in S:
            st = S[name]; st.font.size = Pt(size); st.font.bold = True; st.font.italic = (name == "Heading 2")
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if "Bibliography" in S:
        pf = S["Bibliography"].paragraph_format
        pf.left_indent, pf.first_line_indent = Cm(1.27), Cm(-1.27)
    d.save(str(path))


def finish(path):
    """Line numbers (continuous), page numbers, table font."""
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document(str(path))
    for sec in d.sections:
        sp = sec._sectPr
        ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
        ln.set(qn("w:distance"), "360")
        # schema order: lnNumType must come before pgNumType/cols/docGrid
        anchor = None
        for tag in ("w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                    "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings"):
            anchor = sp.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(ln)
        else:
            sp.append(ln)
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = 1
        ppr = p._p.get_or_add_pPr(); sup = OxmlElement("w:suppressLineNumbers")
        jc = ppr.find(qn("w:jc"))
        (jc.addprevious(sup) if jc is not None else ppr.append(sup))
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r); p._p.append(fld)
    # tables: full text width (16 cm = 9072 dxa), fixed layout, column widths proportional to the longest word/entry
    TOTAL = 9072
    for tbl in d.tables:
        t = tbl._tbl; pr = t.tblPr
        ncol = max(len(r.cells) for r in tbl.rows)
        need = [0] * ncol
        for r in tbl.rows:
            for k, c in enumerate(r.cells[:ncol]):
                txt = c.text.strip()
                longest = max([len(w) for w in txt.split()] + [0])
                need[k] = max(need[k], min(len(txt), 28), longest + 2)
        need = [max(n, 6) for n in need]
        widths = [int(TOTAL * n / sum(need)) for n in need]
        for old in pr.findall(qn("w:tblW")) + pr.findall(qn("w:tblLayout")):
            pr.remove(old)
        w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(TOTAL)); w.set(qn("w:type"), "dxa"); pr.append(w)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
        grid = t.find(qn("w:tblGrid"))
        if grid is not None:
            for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(wd))
        for r in tbl.rows:
            for c, wd in zip(r.cells, widths):
                tcpr = c._tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
                tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(wd))
    d.save(str(path))
    _xml_fixes(path)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
ORDER = {
    W + "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
    "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData embedTrueTypeFonts "
    "embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges bordersDoNotSurroundHeader "
    "bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState "
    "formsDesign attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
    "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride "
    "styleLockTheme styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone "
    "doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders "
    "bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing "
    "displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin "
    "drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning "
    "characterSpacingControl printTwoOnOne strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore "
    "savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText "
    "doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace "
    "updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids mathPr attachedSchema themeFontLang "
    "clrSchemeMapping doNotIncludeSubdocsInStats doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown "
    "smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags decimalSymbol listSeparator",
    W + "style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked "
    "personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
    W + "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
    "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
    "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    W + "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
    "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang "
    "eastAsianLayout specVanish oMath",
    W + "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
    "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    W + "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    W + "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType "
    "pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange",
}
ORDER[M + "rPr"] = "lit nor scr sty brk aln"
ORDER[M + "mcPr"] = "count mcJc"
ORDER = {k: {t: i for i, t in enumerate(v.split())} for k, v in ORDER.items()}


def _sort_children(root):
    for el in root.iter():
        rank = ORDER.get(el.tag)
        if rank is None:
            continue
        kids = list(el)
        key = {id(c): rank.get(c.tag.split("}")[-1], len(rank)) for c in kids}
        srt = sorted(kids, key=lambda c: key[id(c)])     # stable: unknown children keep their order at the end
        if srt != kids:
            for c in kids:
                el.remove(c)
            el.extend(srt)


def _xml_fixes(path):
    """Schema fixes: element order (settings, styles, pPr, rPr, tblPr, tcPr, sectPr, m:dPr), m:nor with m:sty,
    8-digit numbering nsids, pgMar header/footer/gutter, theme fonts overriding Times New Roman."""
    import zipfile, shutil
    from lxml import etree
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename in ("word/document.xml", "word/styles.xml", "word/settings.xml", "word/numbering.xml"):
                if item.filename == "word/styles.xml":
                    data = re.sub(rb' w:(ascii|hAnsi|eastAsia|cs)Theme="[^"]*"', b"", data)
                if item.filename == "word/settings.xml":
                    data = re.sub(rb"<w:zoom [^>]*/>", b"", data)
                root = etree.fromstring(data)
                _sort_children(root)
                for dpr in root.iter(M + "dPr"):
                    kids = list(dpr); order = {"begChr": 0, "sepChr": 1, "endChr": 2, "grow": 3, "shp": 4, "ctrlPr": 5}
                    for c in kids:
                        dpr.remove(c)
                    dpr.extend(sorted(kids, key=lambda c: order.get(c.tag.split("}")[-1], 9)))
                for dpr in root.iter(M + "dPr"):            # empty separator with a single argument: drop it
                    sep = dpr.find(M + "sepChr")
                    if sep is not None and sep.get(M + "val") == "" and len(dpr.getparent().findall(M + "e")) <= 1:
                        dpr.remove(sep)
                for mrpr in root.iter(M + "rPr"):
                    if mrpr.find(M + "nor") is not None:
                        for sty in mrpr.findall(M + "sty"):
                            mrpr.remove(sty)
                for nsid in root.iter(W + "nsid"):
                    nsid.set(W + "val", nsid.get(W + "val").zfill(8)[-8:])
                for pm in root.iter(W + "pgMar"):
                    for a, v in (("header", "708"), ("footer", "708"), ("gutter", "0")):
                        if pm.get(W + a) is None:
                            pm.set(W + a, v)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zo.writestr(item, data)
    shutil.move(tmp, path)


def highlights(path, ref):
    md = "# Highlights\n\n" + "\n".join(f"- {h}" for h in HIGHLIGHTS) + "\n"
    for h in HIGHLIGHTS:
        assert len(h) <= 85, (len(h), h)
    subprocess.run(["pandoc", "-f", "markdown", "-o", str(path), f"--reference-doc={ref}"], input=md.encode(), check=True)
    _xml_fixes(path)


def main():
    lab = labels()
    tex = (HERE / "main.tex").read_text()
    tex = flatten_inputs(tex)
    tex = tex.replace("\\linenumbers", "")
    tex, title = title_block(tex)
    tex = resolve_ifexists(tex)
    tex = strip_resizebox(tex)
    tex = convert_floats(tex, lab)
    tex = number_headings(tex)
    tex = resolve_refs(tex, lab)
    # numbered displays show the numbers of the PDF
    def eqnum(m):          # single equations: the number goes at the end of the display
        body = m.group(1)
        lm = re.search(r"\\label\{(eq:[^}]+)\}", body)
        if not lm:
            return m.group(0)
        body = body.replace(lm.group(0), "").rstrip()
        return "\\begin{equation}" + body + "\\qquad(" + lab.get(lm.group(1), ("", ""))[0] + ")\\end{equation}"
    tex = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}", eqnum, tex, flags=re.S)
    tex = re.sub(r"\\label\{(eq:[^}]+)\}", lambda m: "\\qquad(" + lab.get(m.group(1), ("", ""))[0] + ")", tex)
    tex = re.sub(r"\\label\{[^}]+\}", "", tex)
    tex = tex.replace("\\Bigl", "").replace("\\Bigr", "")
    tex = tex.replace("$^{\\star}$", "\\textsuperscript{*}")
    for w in ("mean", "anchor", "sign"):
        tex = tex.replace("\\text{%s:}" % w, "\\text{%s: }" % w)
    tex = re.sub(r"\\bibliographystyle\{[^}]*\}\s*\\bibliography\{[^}]*\}", "", tex)
    tex = tex.replace("\\begin{document}", "\\title{" + title + "}\n\\begin{document}")
    (HERE / "main_flat.tex").write_text(tex)
    ref = HERE / "build_docx" / "reference.docx"; ref.parent.mkdir(exist_ok=True)
    reference_doc(ref)
    r = subprocess.run(["pandoc", "main_flat.tex", "-o", "main.docx", "--citeproc", "--csl=apa.csl",
                        "--bibliography=refs.bib", f"--reference-doc={ref}", "--resource-path=.",
                        "-M", "link-citations=true", "-M", "reference-section-title=References"],
                       cwd=HERE, capture_output=True, text=True)
    warn = [l for l in r.stderr.splitlines() if "WARNING" in l]
    print("pandoc rc", r.returncode, f"{len(warn)} warnings"); print("\n".join(warn[:10]))
    finish(HERE / "main.docx")
    highlights(HERE / "highlights.docx", ref)
    print("wrote main.docx, highlights.docx")


if __name__ == "__main__":
    main()
